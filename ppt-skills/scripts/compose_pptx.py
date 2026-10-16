#!/usr/bin/env python3
"""Compose an editable .pptx from one or more page layout.json files.

Each layout describes one slide as stacked layers (back to front):
  background (full-bleed) -> frame (full-bleed transparent PNG) ->
  icons (positioned PNGs) -> texts (real, editable text boxes)

All boxes use SOURCE-IMAGE PIXELS: ref_width/ref_height is the pixel size of
the slide image the boxes were measured on. Relative file paths resolve
against the layout's own directory, so layouts can be merged from anywhere.

Usage:
    python compose_pptx.py -o page.pptx pages/03/layout.json
    python compose_pptx.py -o final.pptx pages/01/layout.json pages/02/layout.json ...
    python compose_pptx.py -o page.pptx pages/03/layout.json --render pages/03/qa/render

layout.json:
{
  "ref_width": 2400, "ref_height": 1350,
  "background": "background.png",
  "frame": "frame.png",
  "icons": [{"file": "icons/ic_01.png", "x": 120, "y": 260, "w": 96, "h": 96}],
  "texts": [{"text": "标题", "x": 120, "y": 60, "w": 900, "h": 80,
             "size_px": 64, "color": "#18324F", "bold": true,
             "align": "left", "valign": "middle", "line_spacing": 1.0,
             "font": "Microsoft YaHei"}]
}
Text size: "size" (pt) wins, else "size_px" (character height in source px).
Mixed styles in one box: "runs": [{"text": "...", "color": "#C00000", "bold": true}, ...];
"\\n" inside text or runs starts a new paragraph.
Optional slide size: "slide_width_in"/"slide_height_in" (default width 13.333in,
height from the source aspect ratio).
"""
from __future__ import annotations

import argparse
import json
import subprocess
import sys
from pathlib import Path

EMU_PER_INCH = 914400
DEFAULT_WIDTH_IN = 13.333
SCRIPTS_DIR = Path(__file__).resolve().parent


def die(msg: str):
    print(f"Error: {msg}", file=sys.stderr)
    raise SystemExit(2)


def load_layout(path: Path) -> dict:
    try:
        data = json.loads(Path(path).read_text(encoding="utf-8-sig"))
    except (OSError, json.JSONDecodeError) as e:
        die(f"cannot read layout {path}: {e}")
    if not isinstance(data, dict):
        die(f"{path}: layout must be a JSON object")
    for key in ("ref_width", "ref_height"):
        if not isinstance(data.get(key), (int, float)) or data[key] <= 0:
            die(f"{path}: {key} must be the source image pixel size")
    data["_dir"] = Path(path).resolve().parent
    return data


def slide_size_in(layout: dict) -> tuple[float, float]:
    w = float(layout.get("slide_width_in") or DEFAULT_WIDTH_IN)
    h = layout.get("slide_height_in")
    h = float(h) if h else w * layout["ref_height"] / layout["ref_width"]
    return w, h


def resolve(layout: dict, file: str) -> Path:
    p = Path(file)
    return p if p.is_absolute() else layout["_dir"] / p


def text_size_pt(item: dict, layout: dict, slide_h_pt: float, default: float | None = None) -> float:
    if item.get("size") is not None:
        return float(item["size"])
    if item.get("size_px") is not None:
        return float(item["size_px"]) * slide_h_pt / float(layout["ref_height"])
    return default if default is not None else 18.0


def _rgb(value: str):
    from pptx.dml.color import RGBColor
    v = str(value).strip().lstrip("#")
    if len(v) == 3:
        v = "".join(c * 2 for c in v)
    if len(v) != 6:
        return RGBColor(0x11, 0x11, 0x11)
    return RGBColor(int(v[0:2], 16), int(v[2:4], 16), int(v[4:6], 16))


def _set_fonts(run, name: str):
    """Set latin + east-asian + complex-script typefaces so CJK renders as requested."""
    from pptx.oxml.ns import qn
    run.font.name = name
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set("typeface", name)


def _paragraphs(item: dict):
    """Split text/runs into paragraphs: list of lists of (text, run_style_dict)."""
    runs = item.get("runs") or [{"text": item.get("text", "")}]
    paras = [[]]
    for r in runs:
        parts = str(r.get("text", "")).split("\n")
        for i, part in enumerate(parts):
            if i > 0:
                paras.append([])
            if part:
                paras[-1].append((part, r))
    return paras


def add_slide(prs, layout: dict, sw_emu: int, sh_emu: int, slide_h_pt: float, label: str):
    from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
    from pptx.util import Emu, Pt

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    kx = sw_emu / float(layout["ref_width"])
    ky = sh_emu / float(layout["ref_height"])

    def box(item):
        try:
            x, y, w, h = (float(item[k]) for k in ("x", "y", "w", "h"))
        except (KeyError, TypeError, ValueError):
            die(f"{label}: every icon/text needs numeric x/y/w/h: {item}")
        return Emu(round(x * kx)), Emu(round(y * ky)), Emu(max(1, round(w * kx))), Emu(max(1, round(h * ky)))

    for key in ("background", "frame"):
        if layout.get(key):
            p = resolve(layout, layout[key])
            if not p.exists():
                die(f"{label}: {key} not found: {p}")
            slide.shapes.add_picture(str(p), 0, 0, width=Emu(sw_emu), height=Emu(sh_emu))

    for ic in layout.get("icons", []):
        p = resolve(layout, ic.get("file", ""))
        if not p.is_file():
            die(f"{label}: icon not found: {p}")
        x, y, w, h = box(ic)
        slide.shapes.add_picture(str(p), x, y, width=w, height=h)

    align_map = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER,
                 "right": PP_ALIGN.RIGHT, "justify": PP_ALIGN.JUSTIFY}
    anchor_map = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE,
                  "center": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}

    for tx in layout.get("texts", []):
        x, y, w, h = box(tx)
        tb = slide.shapes.add_textbox(x, y, w, h)
        tf = tb.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.vertical_anchor = anchor_map.get(tx.get("valign", "top"), MSO_ANCHOR.TOP)
        # Boxes are measured around the visible text, so default insets would shift it.
        tf.margin_left = Pt(float(tx.get("margin_left", 0)))
        tf.margin_right = Pt(float(tx.get("margin_right", 0)))
        tf.margin_top = Pt(float(tx.get("margin_top", 0)))
        tf.margin_bottom = Pt(float(tx.get("margin_bottom", 0)))

        base_pt = text_size_pt(tx, layout, slide_h_pt)
        for pi, runs in enumerate(_paragraphs(tx)):
            para = tf.paragraphs[0] if pi == 0 else tf.add_paragraph()
            para.alignment = align_map.get(tx.get("align", "left"), PP_ALIGN.LEFT)
            if tx.get("line_spacing"):
                para.line_spacing = float(tx["line_spacing"])
            for text, style in runs:
                run = para.add_run()
                run.text = text
                run.font.size = Pt(text_size_pt(style, layout, slide_h_pt, default=base_pt))
                run.font.bold = bool(style.get("bold", tx.get("bold", False)))
                run.font.italic = bool(style.get("italic", tx.get("italic", False)))
                run.font.color.rgb = _rgb(style.get("color", tx.get("color", "#111111")))
                _set_fonts(run, str(style.get("font", tx.get("font", "Microsoft YaHei"))))


def build_pptx(layout_paths: list[Path], out_path: Path) -> Path:
    from pptx import Presentation
    from pptx.util import Emu

    layouts = [load_layout(Path(p)) for p in layout_paths]
    if not layouts:
        die("no layouts given")
    sw_in, sh_in = slide_size_in(layouts[0])
    for path, lay in zip(layout_paths, layouts):
        w, h = slide_size_in(lay)
        if abs(w / h - sw_in / sh_in) > 0.01 * sw_in / sh_in:
            print(f"Warning: {path} aspect ratio differs from the first page; it will be stretched.",
                  file=sys.stderr)

    prs = Presentation()
    sw_emu, sh_emu = round(sw_in * EMU_PER_INCH), round(sh_in * EMU_PER_INCH)
    prs.slide_width, prs.slide_height = Emu(sw_emu), Emu(sh_emu)
    for i, (path, lay) in enumerate(zip(layout_paths, layouts), 1):
        add_slide(prs, lay, sw_emu, sh_emu, sh_in * 72.0, f"slide {i} ({path})")

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    print(f"Wrote {out_path} ({len(layouts)} slides)")
    return out_path


def render_with_powerpoint(pptx: Path, out_dir: Path, width: int = 1920) -> list[Path]:
    """Render every slide of a real .pptx through Microsoft PowerPoint (Windows only)."""
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    for old in out_dir.glob("slide_*.png"):
        old.unlink()
    cmd = ["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass",
           "-File", str(SCRIPTS_DIR / "render_pptx.ps1"),
           "-InputPptx", str(Path(pptx).resolve()), "-OutputDir", str(out_dir.resolve()),
           "-Width", str(width)]
    try:
        res = subprocess.run(cmd, capture_output=True, text=True, timeout=900)
    except (OSError, subprocess.TimeoutExpired) as e:
        die(f"PowerPoint render failed to run: {e}")
    if res.returncode != 0:
        die(f"PowerPoint render failed:\n{res.stdout}\n{res.stderr}")
    pngs = sorted(out_dir.glob("slide_*.png"))
    if not pngs:
        die(f"PowerPoint render produced no PNGs in {out_dir}")
    print(f"Rendered {len(pngs)} slide(s) with PowerPoint -> {out_dir}")
    return pngs


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("layouts", nargs="+", help="Page layout.json files, in slide order.")
    ap.add_argument("-o", "--out", required=True, help="Output .pptx path.")
    ap.add_argument("--render", metavar="DIR", help="Also render the result with PowerPoint into DIR.")
    args = ap.parse_args()
    out = build_pptx([Path(p) for p in args.layouts], Path(args.out))
    if args.render:
        render_with_powerpoint(out, Path(args.render))


if __name__ == "__main__":
    main()
