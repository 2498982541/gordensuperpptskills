"""Script tests with synthetic assets; no image generation.

Run: python -m unittest discover -s tests -v
Set PPT_SKILLS_OFFICE_TESTS=1 on Windows with PowerPoint to also test real rendering.
"""
from __future__ import annotations

import json
import os
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

SCRIPTS = Path(__file__).resolve().parents[1] / "ppt-skills" / "scripts"
sys.path.insert(0, str(SCRIPTS))
from compose_pptx import build_pptx  # noqa: E402
from layout_guard import check_layout, fix_scale  # noqa: E402


def run(script: str, *args: str, env: dict | None = None) -> subprocess.CompletedProcess:
    return subprocess.run([sys.executable, str(SCRIPTS / script), *map(str, args)],
                          capture_output=True, text=True, encoding="utf-8", env=env)


def make_page(page: Path, size=(1600, 900)) -> Path:
    """A fake converted page: source, background, frame, one icon, layout.json."""
    w, h = size
    (page / "icons").mkdir(parents=True)
    src = Image.new("RGB", size, "#F4F6F9")
    d = ImageDraw.Draw(src)
    d.rectangle([w * .1, h * .3, w * .45, h * .8], fill="#FFFFFF", outline="#1B3F6E", width=4)
    src.save(page / "source.png")
    Image.new("RGB", size, "#F4F6F9").save(page / "background.png")
    frame = Image.new("RGBA", size, (0, 0, 0, 0))
    ImageDraw.Draw(frame).rectangle([w * .1, h * .3, w * .45, h * .8], fill="#FFFFFF", outline="#1B3F6E", width=4)
    frame.save(page / "frame.png")
    Image.new("RGBA", (80, 80), "#B5121B").save(page / "icons" / "ic_01.png")
    layout = {
        "ref_width": w, "ref_height": h,
        "background": "background.png", "frame": "frame.png",
        "icons": [{"file": "icons/ic_01.png", "x": w * .6, "y": h * .3, "w": h * .1, "h": h * .1}],
        "texts": [
            {"text": "可编辑标题", "x": w * .1, "y": h * .08, "w": w * .6, "h": h * .1,
             "size_px": h * .07, "color": "#1B3F6E", "bold": True},
            {"runs": [{"text": "普通文字"}, {"text": "高亮", "color": "#C00000", "bold": True},
                      {"text": "\n第二段"}],
             "x": w * .12, "y": h * .35, "w": w * .3, "h": h * .08, "size_px": h * .03,
             "line_spacing": 1.2},
        ],
    }
    (page / "layout.json").write_text(json.dumps(layout, ensure_ascii=False), encoding="utf-8")
    return page / "layout.json"


class ComposeTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def test_single_page_layers_and_text(self):
        layout = make_page(self.root / "pages" / "01")
        out = build_pptx([layout], self.root / "page.pptx")
        prs = Presentation(str(out))
        self.assertAlmostEqual(prs.slide_width / 914400, 13.333, places=2)
        self.assertAlmostEqual(prs.slide_height / 914400, 7.5, places=2)
        shapes = list(prs.slides[0].shapes)
        pictures = [s for s in shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
        boxes = [s for s in shapes if s.has_text_frame]
        self.assertEqual(len(pictures), 3)            # background, frame, icon
        self.assertEqual(pictures[0].width, prs.slide_width)
        self.assertEqual(boxes[0].text_frame.text, "可编辑标题")
        run = boxes[0].text_frame.paragraphs[0].runs[0]
        self.assertAlmostEqual(run.font.size.pt, 0.07 * 540, places=0)   # size_px -> pt
        self.assertTrue(run.font.bold)
        self.assertEqual(run._r.get_or_add_rPr().find(qn("a:ea")).get("typeface"), "Microsoft YaHei")
        paras = boxes[1].text_frame.paragraphs
        self.assertEqual([p.text for p in paras], ["普通文字高亮", "第二段"])
        self.assertEqual(str(paras[0].runs[1].font.color.rgb), "C00000")
        self.assertFalse(paras[0].runs[0].font.bold)

    def test_merge_pages_with_different_resolutions(self):
        a = make_page(self.root / "pages" / "01", (1600, 900))
        b = make_page(self.root / "pages" / "02", (2400, 1350))
        out = build_pptx([a, b], self.root / "final.pptx")
        prs = Presentation(str(out))
        self.assertEqual(len(prs.slides), 2)
        ic1 = [s for s in prs.slides[0].shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE][2]
        ic2 = [s for s in prs.slides[1].shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE][2]
        self.assertAlmostEqual(ic1.left, ic2.left, delta=2)   # same relative position

    def test_cli_resolves_assets_from_layout_dir(self):
        layout = make_page(self.root / "pages" / "01")
        res = subprocess.run([sys.executable, str(SCRIPTS / "compose_pptx.py"), "-o", "out.pptx", str(layout)],
                             cwd=self.root, capture_output=True, text=True)
        self.assertEqual(res.returncode, 0, res.stderr)
        self.assertTrue((self.root / "out.pptx").exists())


class GuardTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.page = self.tmp_page = Path(self.tmp.name) / "page"
        self.layout = make_page(self.page)

    def tearDown(self):
        self.tmp.cleanup()

    def edit(self, fn):
        lay = json.loads(self.layout.read_text(encoding="utf-8"))
        fn(lay)
        self.layout.write_text(json.dumps(lay, ensure_ascii=False), encoding="utf-8")

    def test_clean_layout_passes(self):
        self.assertEqual(check_layout(self.page / "source.png", self.layout), ([], []))

    def test_half_scale_measurement_is_caught_and_fixed(self):
        def halve(lay):
            lay["ref_width"], lay["ref_height"] = 800, 450
            for item in lay["icons"] + lay["texts"]:
                for k in ("x", "y", "w", "h", "size_px"):
                    if k in item:
                        item[k] /= 2
        self.edit(halve)
        _, errors = check_layout(self.page / "source.png", self.layout)
        self.assertTrue(any("--fix-scale" in e for e in errors))
        fix_scale(self.page / "source.png", self.layout)
        self.assertEqual(check_layout(self.page / "source.png", self.layout), ([], []))
        lay = json.loads(self.layout.read_text(encoding="utf-8"))
        self.assertEqual(lay["texts"][0]["size_px"], 63.0)

    def test_out_of_canvas_and_missing_file(self):
        self.edit(lambda lay: lay["icons"].append({"file": "icons/nope.png", "x": 1590, "y": 10, "w": 50, "h": 50}))
        _, errors = check_layout(self.page / "source.png", self.layout)
        self.assertTrue(any("outside" in e for e in errors))
        self.assertTrue(any("not found" in e for e in errors))

    def test_tiny_font_and_all_bold_warn(self):
        def style(lay):
            lay["texts"][0]["size_px"] = 5
            lay["texts"] += [{"text": f"正文{i}", "x": 100, "y": 600 + i * 30, "w": 300, "h": 28,
                              "size_px": 24, "bold": True} for i in range(6)]
            lay["texts"][1]["bold"] = True
        self.edit(style)
        warnings, _ = check_layout(self.page / "source.png", self.layout)
        self.assertTrue(any("only" in w and "pt" in w for w in warnings))
        self.assertTrue(any("bold" in w for w in warnings))


class ToolTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def test_build_page_layers_and_boxes(self):
        page = self.root / "page"
        make_page(page)
        for stage, out in (("layers", "layers.png"), ("boxes", "source_boxes.png")):
            res = run("build_page.py", page, "--stage", stage)
            self.assertEqual(res.returncode, 0, res.stdout + res.stderr)
            self.assertTrue((page / "qa" / out).exists())

    def test_ruler_region(self):
        page = self.root / "page"
        make_page(page)
        res = run("ruler.py", page / "source.png", self.root / "zoom.png", "--region", 100, 200, 400, 225)
        self.assertEqual(res.returncode, 0, res.stderr)
        with Image.open(self.root / "zoom.png") as im:
            self.assertEqual(im.size, (1800, 1012))

    def test_chroma_key_and_slice(self):
        sheet = Image.new("RGB", (800, 800), "#00ff00")
        d = ImageDraw.Draw(sheet)
        d.ellipse([80, 80, 320, 320], fill="#B5121B")
        d.rectangle([480, 480, 720, 720], fill="#1B3F6E")
        sheet.save(self.root / "icons_raw.png")
        res = run("chroma_key.py", "--input", self.root / "icons_raw.png", "--out", self.root / "icons_t.png",
                  "--preset", "icon-safe", "--force")
        self.assertEqual(res.returncode, 0, res.stderr)
        res = run("slice_grid.py", self.root / "icons_t.png", self.root / "icons", "--auto", "--pad", "24",
                  "--contact-sheet", "--prefix", "ic")
        self.assertEqual(res.returncode, 0, res.stderr)
        manifest = json.loads((self.root / "icons" / "icons_manifest.json").read_text(encoding="utf-8"))
        self.assertEqual(len(manifest["icons"]), 2)
        with Image.open(self.root / "icons" / manifest["icons"][0]["file"]) as icon:
            self.assertEqual(icon.mode, "RGBA")
            self.assertEqual(icon.getpixel((0, 0))[3], 0)            # transparent corner
            r, g, b, a = icon.getpixel((icon.width // 2, icon.height // 2))
            self.assertEqual((a, abs(r - 0xB5) < 6, g < 30), (255, True, True))   # color kept

    def test_check_models_reports_and_acks_new_model(self):
        home = self.root / "codex"
        home.mkdir()
        (home / "models_cache.json").write_text(json.dumps({"models": [
            {"slug": "gpt-6.1-sol", "visibility": "list"},
            {"slug": "gpt-7-sol", "visibility": "list", "description": "new"},
            {"slug": "hidden", "visibility": "hide"}]}), encoding="utf-8")
        config = self.root / "models.json"
        config.write_text(json.dumps({"conversion": {"model": "gpt-6.1-sol", "reasoning_effort": "high"},
                                      "known_models": ["gpt-6.1-sol"]}), encoding="utf-8")
        env = {**os.environ, "CODEX_HOME": str(home), "PYTHONIOENCODING": "utf-8"}
        res = run("check_models.py", "--config", config, env=env)
        self.assertIn("gpt-7-sol", res.stdout)
        self.assertNotIn("hidden", res.stdout)
        run("check_models.py", "--config", config, "--ack", env=env)
        self.assertIn("没有新模型", run("check_models.py", "--config", config, env=env).stdout)
        self.assertEqual(json.loads(config.read_text(encoding="utf-8"))["conversion"]["model"], "gpt-6.1-sol")


@unittest.skipUnless(os.environ.get("PPT_SKILLS_OFFICE_TESTS") == "1", "needs Windows + PowerPoint")
class PowerPointRenderTests(unittest.TestCase):
    def test_build_page_full_renders_with_powerpoint(self):
        with tempfile.TemporaryDirectory() as tmp:
            page = Path(tmp) / "page"
            make_page(page)
            res = run("build_page.py", page)
            self.assertEqual(res.returncode, 0, res.stdout + res.stderr)
            with Image.open(page / "qa" / "render" / "slide_01.png") as im:
                self.assertEqual(im.size, (1920, 1080))
                self.assertGreater(min(im.convert("RGB").getpixel((300, 700))), 240)   # white card from the frame
            for name in ("side_by_side.png", "preview_boxes.png", "blend.png", "diff_heatmap.png"):
                self.assertTrue((page / "qa" / name).exists(), name)


if __name__ == "__main__":
    unittest.main()
